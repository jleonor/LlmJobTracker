from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from localjobtracker.app import TOKEN_HEADER, create_app
from localjobtracker.extraction import dehyphenate, extract
from localjobtracker.runner import FakeRunner

FIXTURES = Path(__file__).parent.parent.parent / "tests" / "fixtures" / "synthetic" / "profile"
LISTING_FIXTURES = (
    Path(__file__).parent.parent.parent / "tests" / "fixtures" / "synthetic" / "job_listing"
)


def test_pdf_with_a_text_layer_is_read() -> None:
    result = extract(FIXTURES / "cv_en.pdf")
    assert result.status == "ok"
    assert result.message is None
    assert "Analyst, Example Corp, 2020 to 2023" in result.text


def test_pdf_without_a_text_layer_is_reported_not_read() -> None:
    result = extract(FIXTURES / "scanned_no_text.pdf")
    assert result.status == "no_text"
    assert result.text == ""
    assert "no text layer" in (result.message or "")


def test_dehyphenate_rejoins_a_line_wrapped_word() -> None:
    """Live bug (2026-10-11): a real PDF's own line wrap split "scikit-learn" into
    "scikit-\\nlearn", which then failed generate_job_details' verbatim-match check against the
    model's natural, unbroken "scikit-learn"."""
    text = "ML frameworks (TensorFlow, PyTorch, scikit-\nlearn).\nCloud Experience: AWS."
    expected = "ML frameworks (TensorFlow, PyTorch, scikit-learn).\nCloud Experience: AWS."
    assert dehyphenate(text) == expected


def test_dehyphenate_absorbs_trailing_and_leading_line_whitespace() -> None:
    assert dehyphenate("scikit-  \n  learn") == "scikit-learn"


def test_dehyphenate_leaves_an_ordinary_line_break_alone() -> None:
    text = "Requirements:\n- Strong Python\n- Strong SQL"
    assert dehyphenate(text) == text


def test_docx_gives_paragraphs_and_table_cells() -> None:
    result = extract(FIXTURES / "cv_en.docx")
    assert result.status == "ok"
    assert "Analyst, Example Corp, 2020 to 2023" in result.text
    assert "Skill | Level" in result.text
    assert "SQL | Advanced" in result.text


def test_plain_text_is_read_as_utf8() -> None:
    result = extract(FIXTURES / "questionnaire.txt")
    assert result.status == "ok"
    assert result.text.startswith("Synthetic questionnaire.")


def test_older_french_text_saved_as_windows_1252_is_read() -> None:
    result = extract(FIXTURES / "questionnaire_fr_cp1252.txt")
    assert result.status == "ok"
    assert "synthétique" in result.text
    assert "Quels postes vous intéressent" in result.text


def test_damaged_pdf_fails_with_a_message(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"this is not a pdf")
    result = extract(broken)
    assert result.status == "failed"
    assert result.text == ""
    assert "Could not read this file" in (result.message or "")


def test_unsupported_extension_is_reported(tmp_path: Path) -> None:
    image = tmp_path / "scan.png"
    image.write_bytes(b"\x89PNG")
    assert extract(image).status == "unsupported"


def test_html_listing_drops_navigation_and_keeps_the_posting() -> None:
    result = extract(LISTING_FIXTURES / "listing.html")
    assert result.status == "ok"
    assert result.message is None
    assert "Senior Data Analyst" in result.text
    assert "4+ years in an analytics or data science role" in result.text
    assert "Accept all" not in result.text
    assert "Subscribe to our newsletter" not in result.text


def test_html_with_only_navigation_and_boilerplate_is_reported_not_read() -> None:
    result = extract(LISTING_FIXTURES / "junk_only.html")
    assert result.status == "no_text"
    assert result.text == ""
    assert "could be found" in (result.message or "")


def test_pdf_listing_fixture_is_read() -> None:
    result = extract(LISTING_FIXTURES / "listing.pdf")
    assert result.status == "ok"
    assert "Senior Data Analyst" in result.text
    assert "Requirements" in result.text


def test_docx_listing_fixture_is_read() -> None:
    result = extract(LISTING_FIXTURES / "listing.docx")
    assert result.status == "ok"
    assert "Senior Data Analyst" in result.text
    assert "Strong SQL and at least one scripting language" in result.text


def test_txt_listing_fixture_is_read() -> None:
    result = extract(LISTING_FIXTURES / "listing.txt")
    assert result.status == "ok"
    assert "Senior Data Analyst" in result.text


@pytest.fixture
def upload_client(db_path: Path) -> Iterator[TestClient]:
    with TestClient(create_app("test-token", db_path, FakeRunner())) as test_client:
        test_client.headers[TOKEN_HEADER] = "test-token"
        yield test_client


def test_upload_records_the_extraction_and_serves_the_text(upload_client: TestClient) -> None:
    content = (FIXTURES / "cv_en.pdf").read_bytes()
    created = upload_client.post(
        "/profile/files/master_cv",
        content=content,
        headers={"X-File-Name": "cv_en.pdf"},
    )
    assert created.status_code == 201
    body = created.json()
    assert body["extraction_status"] == "ok"

    text = upload_client.get(f"/profile/files/{body['id']}/text").json()
    assert text["extraction_status"] == "ok"
    assert "Analyst, Example Corp" in text["text"]


def test_upload_of_a_scan_is_kept_and_reported(upload_client: TestClient) -> None:
    content = (FIXTURES / "scanned_no_text.pdf").read_bytes()
    created = upload_client.post(
        "/profile/files/older_cv",
        content=content,
        headers={"X-File-Name": "scan.pdf"},
    )
    body = created.json()
    assert body["extraction_status"] == "no_text"
    assert upload_client.get("/profile/files").json()[0]["id"] == body["id"]


def test_text_route_returns_404_for_an_unknown_file(upload_client: TestClient) -> None:
    assert upload_client.get("/profile/files/999/text").status_code == 404


def test_a_docx_text_box_and_footer_are_read(tmp_path: Path) -> None:
    import docx as python_docx
    from lxml import etree

    document = python_docx.Document()
    document.add_paragraph("Alex Sample, Analyst")
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Available since 2021"
    box = etree.fromstring(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:txbxContent><w:p><w:r><w:t>Mar 2023 to present</w:t></w:r></w:p></w:txbxContent>"
        "</w:r>"
    )
    document.element.body.append(box)
    path = tmp_path / "older.docx"
    document.save(str(path))

    result = extract(path)

    assert result.status == "ok"
    assert "Alex Sample, Analyst" in result.text
    assert "Mar 2023 to present" in result.text
    assert "Available since 2021" in result.text
