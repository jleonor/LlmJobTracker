"""Tests for docx_editor.py (step 35, job-llm-steps.md §5): reading a .docx into numbered
blocks, and applying EDIT/INSERT AFTER/REMOVE/ADD ROW AFTER by block id.
"""

from pathlib import Path

import docx
import pytest
from docx.document import Document
from docx.enum.text import WD_BREAK

from localjobtracker import docx_editor as de

FIXTURES = (
    Path(__file__).parent.parent.parent / "tests" / "fixtures" / "synthetic" / "job_documents"
)


def load() -> Document:
    return de.read_document((FIXTURES / "master_cv.docx").read_bytes())


def test_list_blocks_numbers_paragraphs_and_table_rows_in_order() -> None:
    blocks = de.list_blocks(load())
    assert [b.id for b in blocks] == [
        "B1",
        "B2",
        "B3",
        "B4",
        "B5",
        "B6",
        "B7",
        "B8",
        "B9",
        "B10",
        "T1R1",
        "T1R2",
    ]
    assert blocks[1] == de.Block(id="B2", kind="paragraph", text="Full-Stack Engineer")
    assert blocks[-1] == de.Block(id="T1R2", kind="row", text="SQL | Advanced")


def test_blocks_as_text_shows_bracketed_ids() -> None:
    text = de.blocks_as_text(de.list_blocks(load()))
    assert "[B1] Alex Sample" in text
    assert "[T1R2] SQL | Advanced" in text


def test_edit_replaces_text_and_keeps_the_paragraph_style() -> None:
    document = load()
    de.apply_operations(document, [de.EditOp(id="B2", text="Marketing Automation Engineer")])
    paragraph = document.paragraphs[1]
    assert paragraph.text == "Marketing Automation Engineer"
    assert paragraph.style is not None
    assert paragraph.style.name == "Normal"


def test_insert_after_clones_the_anchors_style() -> None:
    document = load()
    de.apply_operations(
        document, [de.InsertAfterOp(id="B6", text="Led a dashboard rebuild, cutting time in half")]
    )
    texts = [p.text for p in document.paragraphs]
    styles = {p.text: p.style.name for p in document.paragraphs if p.style is not None}
    assert "Led a dashboard rebuild, cutting time in half" in texts
    index = texts.index("Led a dashboard rebuild, cutting time in half")
    assert texts[index - 1] == "Built reporting dashboards for the sales team"
    assert styles["Led a dashboard rebuild, cutting time in half"] == "List Bullet"


def test_remove_deletes_the_paragraph() -> None:
    document = load()
    de.apply_operations(document, [de.RemoveOp(id="B7")])
    texts = [p.text for p in document.paragraphs]
    assert "Maintained the CRM data pipeline" not in texts
    assert len(texts) == 9  # one fewer than the fixture's 10


def test_add_row_after_clones_the_table_row() -> None:
    document = load()
    de.apply_operations(document, [de.AddRowAfterOp(id="T1R2", cells=["Python", "Intermediate"])])
    table = document.tables[0]
    assert [c.text for c in table.rows[-1].cells] == ["Python", "Intermediate"]
    assert len(table.rows) == 3


def test_all_ids_resolve_against_the_original_document_not_mid_edit_state() -> None:
    """Multiple operations all reference the snapshot Claude actually saw (job-llm-steps.md §5) -
    an earlier insertion in the same batch must not shift what a later operation's id means."""
    document = load()
    de.apply_operations(
        document,
        [
            de.InsertAfterOp(id="B3", text="A new summary line."),
            de.EditOp(id="B5", text="Senior Analyst, Acme Corp, 2020 to 2024"),
        ],
    )
    texts = [p.text for p in document.paragraphs]
    assert "A new summary line." in texts
    assert "Senior Analyst, Acme Corp, 2020 to 2024" in texts


def test_apply_rejects_an_unknown_block_id_without_changing_the_document() -> None:
    document = load()
    before = [p.text for p in document.paragraphs]
    with pytest.raises(de.DocxEditError, match="Unknown block id"):
        de.apply_operations(document, [de.EditOp(id="B99", text="x")])
    assert [p.text for p in document.paragraphs] == before


def test_apply_rejects_editing_a_table_row_directly() -> None:
    document = load()
    with pytest.raises(de.DocxEditError, match="EDIT only applies to paragraphs"):
        de.apply_operations(document, [de.EditOp(id="T1R2", text="SQL | Expert")])


def test_apply_rejects_add_row_after_on_a_paragraph_id() -> None:
    document = load()
    with pytest.raises(de.DocxEditError, match="ADD ROW AFTER needs a table row id"):
        de.apply_operations(document, [de.AddRowAfterOp(id="B2", cells=["x", "y"])])


def test_apply_rejects_a_cell_count_mismatch() -> None:
    document = load()
    with pytest.raises(de.DocxEditError, match="2 columns"):
        de.apply_operations(document, [de.AddRowAfterOp(id="T1R2", cells=["only one"])])


def test_apply_rejects_no_operations() -> None:
    with pytest.raises(de.DocxEditError, match="No edits"):
        de.apply_operations(load(), [])


def test_document_bytes_round_trips_through_a_fresh_read() -> None:
    document = load()
    de.apply_operations(document, [de.EditOp(id="B2", text="Platform Engineer")])
    content = de.document_bytes(document)
    reopened = de.read_document(content)
    assert reopened.paragraphs[1].text == "Platform Engineer"


# --- strip_manual_page_breaks (D-164) ----------------------------------------------------------
#
# Live bug (2026-10-12): a real master CV had a manual page break on a blank spacer paragraph, to
# land its older-roles section on page 2. Tailoring legitimately adds a line or two ahead of it
# (a summary sentence, an extra bullet), which the break's fixed position does not account for -
# confirmed against the real file that this pushed a handful of lines onto an almost-empty extra
# page, with the rest of the document pushed one page further still.


def test_strip_manual_page_breaks_removes_a_run_level_page_break() -> None:
    document = docx.Document()
    document.add_paragraph("Before the break.")
    spacer = document.add_paragraph()
    spacer.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("After the break.")

    de.strip_manual_page_breaks(document)

    content = de.document_bytes(document)
    reopened = de.read_document(content)
    texts = [p.text for p in reopened.paragraphs]
    assert texts == ["Before the break.", "", "After the break."]
    for paragraph in reopened.paragraphs:
        for run in paragraph.runs:
            assert "w:br" not in run._r.xml or 'w:type="page"' not in run._r.xml


def test_strip_manual_page_breaks_removes_a_paragraph_level_page_break_before() -> None:
    document = docx.Document()
    document.add_paragraph("Before the break.")
    after = document.add_paragraph("After the break.")
    after.paragraph_format.page_break_before = True

    de.strip_manual_page_breaks(document)

    content = de.document_bytes(document)
    reopened = de.read_document(content)
    assert reopened.paragraphs[1].paragraph_format.page_break_before is None


def test_strip_manual_page_breaks_leaves_ordinary_text_alone() -> None:
    document = load()
    before = [(p.text, p.style.name if p.style else None) for p in document.paragraphs]
    de.strip_manual_page_breaks(document)
    after = [(p.text, p.style.name if p.style else None) for p in document.paragraphs]
    assert before == after
